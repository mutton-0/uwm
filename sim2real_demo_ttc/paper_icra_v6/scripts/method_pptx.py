"""把方法总览图导出成可编辑的 PPTX：图片是图片，方框/文字/箭头都是 PowerPoint 原生形状，可直接改。
三张编辑图先从 fig_method 的数据源重新裁出来存 PNG，再摆进幻灯片。
输出 paper_icra_v5/figures/method_editable.pptx"""
import json,glob,sys,zlib,os,numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt,Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
sys.path.insert(0,"/home/boyuewang/120/uwm/sim2real_demo_ttc/scripts")
from appearance_transform import transform
R5="/home/boyuewang/120/uwm/sim2real_demo_ttc/results_5090"; P=f"{R5}/paper_icra_v4"; V5=f"{R5}/paper_icra_v5"
UID="A_0095"; MM="ddv2"; OUT=f"{V5}/figures"
MAN=json.load(open(f"{R5}/risk_card_manifest.json")); IDX={x["uid"]:x for k in ("A","B") for x in MAN[k]}
DV={o["uid"]:o for o in json.load(open(f"{P}/det_validate.json"))}
x=IDX[UID]; y0,y1=330,740
a0=np.asarray(Image.open(x["img"]).convert("RGB")); r0=np.asarray(Image.open(f"/data/dataset/risk_card/{x['rm_name']}").convert("RGB"))
n0=transform(a0,kind="night",scope="global",seed=zlib.crc32(UID.encode())%(2**31))
paths=[]
for nm,im in (("frame_orig",a0),("frame_rm",r0),("frame_night",n0)):
    p=f"{OUT}/{nm}.png"; Image.fromarray(im[y0:y1]).save(p,quality=92); paths.append(p)
# (b) 三条规划单独存一张透明底 PNG（轨迹本身难用 PPT 形状复现，单独作为图片层）
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
def load(pat):
    o={}
    for f in sorted(glob.glob(pat)):
        for r in json.load(open(f)):
            if "err" not in r: o[r["uid"]]=r
    return o
C=({**load(f"{R5}/card_{MM}.json"),**load(f"{R5}/card_{MM}_*.json")})[UID]
N=({**load(f"{R5}/card_night_{MM}.json"),**load(f"{R5}/card_night_{MM}_*.json")}).get(UID)
CC={"clean":"#2a5db0","rm":"#c0392b","night":"#8a8a84"}
fig,b=plt.subplots(figsize=(1.5,2.3)); b.axvspan(-1,1,color="#eef1f6",lw=0,zorder=0)
for key,lw in (("clean",2.0),("rm",1.6)):
    w=np.vstack([[0,0],np.asarray(C["actual"][key])]); b.plot(-w[:,1],w[:,0],color=CC[key],lw=lw)
if N:
    w=np.vstack([[0,0],np.asarray(N["night"])]); b.plot(-w[:,1],w[:,0],color=CC["night"],lw=1.6)
cen=np.mean(np.asarray(x["corners_ego"])[:,:2],0)
b.plot(-cen[1],cen[0],marker="*",ms=12,color="#e34948",mec="white",mew=0.6)
b.plot(0,0,marker="^",ms=8,color="#0b0b0b")
b.set_xlim(-4.2,4.2); b.set_ylim(-1.5,21); b.set_xticks([]); b.set_yticks([]); b.axis("off")
plans=f"{OUT}/method_plans.png"; fig.savefig(plans,dpi=300,transparent=True,bbox_inches="tight"); plt.close(fig)
# ---------- 幻灯片 ----------
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
sl=prs.slides.add_slide(prs.slide_layouts[6])
GREY=RGBColor(0x52,0x51,0x4e); INK=RGBColor(0x1b,0x22,0x26)
def tbox(l,t,w,h,txt,sz=14,bold=False,align=PP_ALIGN.LEFT,color=INK):
    s=sl.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=s.text_frame; tf.word_wrap=True
    tf.text=txt; p=tf.paragraphs[0]; p.alignment=align
    for r in p.runs: r.font.size=Pt(sz); r.font.bold=bold; r.font.color.rgb=color; r.font.name="Arial"
    return s
def rbox(l,t,w,h,txt,fill,line,sz=12):
    s=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(l),Inches(t),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill; s.line.color.rgb=line; s.line.width=Pt(0.75)
    s.adjustments[0]=0.08
    tf=s.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; tf.text=txt
    for p in tf.paragraphs:
        p.alignment=PP_ALIGN.CENTER
        for r in p.runs: r.font.size=Pt(sz); r.font.color.rgb=INK; r.font.name="Arial"
    return s
def arrow(x1,y1,x2,y2,w=1.5):
    c=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=GREY; c.line.width=Pt(w)
    return c
tbox(0.35,0.25,6,0.4,"A check-up for driving policies",20,True)
# (a)
tbox(0.35,0.85,4.2,0.35,"(a) one frame, two counterfactual edits",13,True)
LAB=[("logged frame","detector: pedestrian found",RGBColor(0x2a,0x5d,0xb0)),
     ("pedestrian removed","detector: not found",RGBColor(0xc0,0x39,0x2b)),
     ("night rendering","detector: still found",RGBColor(0x8a,0x8a,0x84))]
for i,(p,(l1,l2,col)) in enumerate(zip(paths,LAB)):
    top=1.30+i*1.75
    sl.shapes.add_picture(p,Inches(0.35),Inches(top),width=Inches(4.15))
    t=tbox(0.45,top+0.05,2.2,0.3,l1,11,True,color=RGBColor(0xff,0xff,0xff))
    t.fill.solid(); t.fill.fore_color.rgb=col; t.line.fill.background()
    tbox(0.45,top+1.05,3.0,0.3,l2,10,False,color=RGBColor(0x52,0x51,0x4e))
# (b)
tbox(5.05,0.85,3.4,0.35,"(b) same policy, three queries",13,True)
sl.shapes.add_picture(plans,Inches(5.15),Inches(1.35),height=Inches(3.6))
tbox(4.95,5.15,3.6,0.6,"three plans over the common 2.5 s horizon;\nclearance $c_t$ to the pedestrian's logged future",10,False,PP_ALIGN.CENTER,GREY)
arrow(4.62,3.4,5.02,3.4)
# (c)
tbox(8.85,0.85,4.2,0.35,"(c) four readouts, five exams",13,True)
rbox(8.85,1.45,2.5,1.0,"contact A · clearance C\ntime T · separation S",RGBColor(0xee,0xf1,0xf6),RGBColor(0xc9,0xc8,0xc0),12)
rbox(8.85,2.75,2.5,1.0,"read on P, Q, N\nover 2.5 s",RGBColor(0xf6,0xf1,0xea),RGBColor(0xc9,0xc8,0xc0),12)
for i,t in enumerate(["exposure","hazard sensitivity","scaling","specificity","lighting CFR"]):
    rbox(11.75,1.25+i*0.85,1.45,0.62,t,RGBColor(0xe9,0xef,0xe9),RGBColor(0xb9,0xcb,0xb9),11)
    arrow(11.40,2.60,11.72,1.56+i*0.85,1.0)
arrow(8.45,2.60,8.82,2.60)
prs.save(f"{OUT}/method_editable.pptx"); print("saved",f"{OUT}/method_editable.pptx")
